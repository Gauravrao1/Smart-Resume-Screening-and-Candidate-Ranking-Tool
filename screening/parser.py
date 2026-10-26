"""Extract plain text from uploaded resumes (PDF, DOCX, TXT) safely."""
from __future__ import annotations

import io
import re
import zipfile

MAX_PDF_PAGES = 30
MAX_CHARS = 200_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ParseError(Exception):
    """Raised when a file cannot be read as a resume."""


def _extension(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_CHARS]


def _read_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF"):
        raise ParseError("File is not a valid PDF.")
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ParseError("PDF is password protected.")
            except ParseError:
                raise
            except Exception:
                raise ParseError("PDF is password protected.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n".join((p.extract_text() or "") for p in pages)
    except ParseError:
        raise
    except Exception as exc:  # corrupt file
        raise ParseError(f"Could not read PDF ({type(exc).__name__}).") from exc


def _read_docx(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ParseError("File is not a valid DOCX.")
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Could not read DOCX ({type(exc).__name__}).") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Return cleaned text for a resume file, or raise :class:`ParseError`."""
    ext = _extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ParseError(f"Unsupported file type '{ext or 'none'}'. Use PDF, DOCX or TXT.")
    if not data:
        raise ParseError("File is empty.")
    readers = {".pdf": _read_pdf, ".docx": _read_docx, ".txt": _read_txt}
    return _clean(readers[ext](data))
